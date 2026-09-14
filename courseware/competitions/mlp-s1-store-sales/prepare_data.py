#!/usr/bin/env python3
"""Build and verify the data for the MLP S1 multi-source store-sales challenge.

The challenge is the Data Science Practice module 4 exercise ("Multi-Sources
Sales Prediction"), unchanged: predict quantity_sold for the 409
Neighborhood_Market items from four other stores' files, an API and a scraped
page. Its data was generated once, by
website/public/modules/data-science-practice/module4/exercise/module4_gen-exo.ipynb
(np.random.seed(42), 2000 items P0001..P2000), and the target was never
published. This script recovers it by re-running that generator, and refuses to
write anything unless the regenerated data reproduces what students download.

Public  -> data/CityMart_data.csv            the four files of the train zip,
           data/Greenfield_Grocers_data.csv  byte-identical to the live download
           data/SuperSaver_Outlet_data.xlsx  (sha256-pinned below)
           data/HighStreet_Bazaar_data.json
           data/Neighborhood_Market_data.csv the test features, no target
Private -> data/neighborhood_market_target.csv   item_code,quantity_sold (409 rows)
           data/y_test.csv            the same bytes under the name the platform
                                      requires: it will not start a file_v1 CSV
                                      challenge without an env file called
                                      y_test.csv (backend lifecycle.py), and reads
                                      the upload check's columns and ids from it
           data/full_generated.csv    every column for all 2000 items, before
                                      the generator punched NaNs into the files,
                                      plus weekly_footfall (the database contract)
           data/ladder.json           the MAE ladder printed at the end

**Checks, all asserted — any mismatch stops the script with the difference.**

1. The five public files and the zip match their pinned sha256.
2. The generator is replayed cell by cell, in the notebook's saved execution
   order (cells 0, 11, 14, 15, 16, 17, 18, 19, 20 — execution counts 1..18 are
   sequential, and no cell between 0 and 16 draws from the RNG). The four
   add_null draws only reproduce the published NaNs in that order: CityMart
   package_volume, Greenfield dimension_length, SuperSaver dimension_width,
   HighStreet days_since_last_purchase, three items each.
3. The regenerated per-store outputs equal the downloads: the three CSVs and the
   JSON byte for byte; the xlsx cell by cell (an xlsx carries a creation
   timestamp, so its bytes cannot be regenerated). The xlsx "Info" sheet has a
   blank first header cell and every data row shifted one column left of its
   header — the author's pandas wrote MultiIndex columns that way with
   index=False; current pandas refuses to, so the expected grid is spelled out.
   Train-store quantity_sold is inside those files, so it is checked too.
4. unit_cost equals the live API (/api/exercise/<password>/prices) for all 2000
   items; customer_score and total_reviews equal ScrapableDataExercise.js and
   the records in the live site bundle behind /module4/scrapable-data.
5. The hidden target reproduces the notebook's own printed output: its first
   rows (P0002 164, P0004 248, P0005 163) and cell 6's "final test on E" MAE,
   3.493040250441976.
6. The exercise's reference solution (module4_exercise_model_test-Copy1.ipynb)
   printed three 5-fold CV scores; assembling the public sources the way it did
   reproduces them (45.0406, 44.1014, 40.0277).

**How the target was made (it matters for the ladder).** Train stores' sales
are a linear formula times a store effect (CityMart 1.2, Greenfield 0.7,
SuperSaver 1.1, HighStreet 1.4) plus N(0, 20) noise per item. Neighborhood_
Market's are NOT that formula with effect 1.0: the generator fits one
LinearRegression on the four train stores pooled, predicts Neighborhood_Market,
truncates to int and adds a single N(0, 6) draw (3.76) to every item. So the
target is the pooled model's prediction shifted by +3.76 — an exactly linear
function of the nine clean features, and Neighborhood_Market's effective store
multiplier is about 1.117 (the pooled 1.10 plus the shift) rather than the 1.0
in the generator's table. The database contract therefore gives it a weekly
footfall of 13400, not 12000.

**The ladder** (measured 2026-09-14; the exercise's get_simple_baseline
recipe — fillna -1, drop store_name and last_modified, StandardScaler,
LinearRegression — trained on the four stores, test MAE on all 409 items):

    constant train mean                     23.18   fail (bar: MAE <= 20)
    files only                              19.89   pass   cv 45.04
    + unit_cost (API)                       17.05   pass   cv 44.10
    + customer_score, total_reviews         3.51    pass   cv 40.03
    + weekly_footfall as a feature          2.15    pass   cv 15.97
    footfall as a multiplier (y/footfall)   2.11    pass

The 3.51 is almost entirely the +3.76 shift (mean signed error -3.51), which a
store-level feature can absorb. Sweep of Neighborhood_Market's footfall, MAE
as a feature / as a multiplier: 12000 21.35 / 21.39, 13000 6.71 / 6.39,
13200 3.95 / 3.59, 13400 2.15 / 2.11, 14400 13.79 / 14.65.

**The target is reproducible by anyone.** The generator notebook is public, on
the website and on GitHub, seeded. Re-running it yields this exact file. Nothing
here can close that without changing the exercise's data.

The API password is fetched at run time and never written to disk.

    uv run --with pandas --with scikit-learn --with openpyxl \\
        python competitions/mlp-s1-store-sales/prepare_data.py
"""
import hashlib
import io
import json
import os
import re
import shutil
import urllib.request
import warnings
import zipfile
from datetime import datetime, timedelta

import numpy as np
import openpyxl
import pandas as pd
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_absolute_error
from sklearn.model_selection import KFold
from sklearn.preprocessing import StandardScaler

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))

SITE = "https://www.raphaelcousin.com"
EXERCISE_URL = f"{SITE}/modules/data-science-practice/module4/exercise"
TRAIN_ZIP_URL = f"{EXERCISE_URL}/module4_exercise_train.zip"
TEST_CSV_URL = f"{EXERCISE_URL}/Neighborhood_Market_data.csv"
API_AUTH_URL = f"{SITE}/api/exercise/auth"
API_PRICES_URL = SITE + "/api/exercise/{password}/prices"
SCRAPE_PAGE_URL = f"{SITE}/module4/scrapable-data"
SCRAPE_JS = os.path.join(
    REPO, "website", "src", "pages", "data-science-practice", "module4",
    "ScrapableDataExercise.js")

TRAIN_STORES = ["CityMart", "Greenfield_Grocers", "SuperSaver_Outlet",
                "HighStreet_Bazaar"]
TEST_STORE = "Neighborhood_Market"
STORE_FILES = {
    "CityMart": "CityMart_data.csv",
    "Greenfield_Grocers": "Greenfield_Grocers_data.csv",
    "SuperSaver_Outlet": "SuperSaver_Outlet_data.xlsx",
    "HighStreet_Bazaar": "HighStreet_Bazaar_data.json",
    "Neighborhood_Market": "Neighborhood_Market_data.csv",
}
PUBLIC_FILES = list(STORE_FILES.values())

# sha256 of the live downloads, verified against the generator on 2026-09-14.
TRAIN_ZIP_SHA256 = (
    "bd0ad817b1af4f50cd7d0e094b6dfd1054d073e1bd10201e37a3de23c5e642a3")
PUBLIC_SHA256 = {
    "CityMart_data.csv":
        "6c1a73b7baaa519f292f3295da9594e1e14677616c559166bd3fadea56fc1014",
    "Greenfield_Grocers_data.csv":
        "318f549e3ccb54cfca193c4c4955b56c199d4311acd129a35f82734a3b8f33ff",
    "SuperSaver_Outlet_data.xlsx":
        "680352587195ec20a1b0306ff074cdf3deb3a5be6fdeb788fa190e3784760cb0",
    "HighStreet_Bazaar_data.json":
        "d5324ca054bce9eaf16fe18edfa2188b79e33564b1ed134faa274da6d6eda9e4",
    "Neighborhood_Market_data.csv":
        "f33e7c019646d9a28149a566deeb86f8b165561fdbe907a62bdf89728964316f",
}

# The database contract: retail.stores.weekly_footfall = the store's sales
# multiplier x 12000. Neighborhood_Market's is its effective multiplier
# (~1.117, see the docstring), not the 1.0 in STORE_EFFECT.
STORE_EFFECT = {"CityMart": 1.2, "Greenfield_Grocers": 0.7,
                "SuperSaver_Outlet": 1.1, "HighStreet_Bazaar": 1.4,
                "Neighborhood_Market": 1.0}
WEEKLY_FOOTFALL = {"CityMart": 14400, "Greenfield_Grocers": 8400,
                   "SuperSaver_Outlet": 13200, "HighStreet_Bazaar": 16800,
                   "Neighborhood_Market": 13400}

# Printed by module4_gen-exo.ipynb (cell 0 head, cell 6 last line).
NOTEBOOK_TARGET_HEAD = {"P0002": 164, "P0004": 248, "P0005": 163}
NOTEBOOK_FINAL_TEST_MAE = 3.493040250441976
# Printed by module4_exercise_model_test-Copy1.ipynb (cells 22, 29, 35).
REFERENCE_CV_MAE = {"files": 45.04056766700087,
                    "files+api": 44.10144429622424,
                    "files+api+scraping": 40.02769508950211}

PASS_MAE = 20.0  # ERROR_THRESHOLD in tests/data-science-practice/module4/exercise1.sh

FILE_COLS = ["mass", "dimension_length", "dimension_width",
             "dimension_height", "days_since_last_purchase",
             "package_volume", "stock_age"]


# --------------------------------------------------------------------------- #
# Download
# --------------------------------------------------------------------------- #
def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": "prepare_data"})
    with urllib.request.urlopen(req, timeout=60) as resp:
        if resp.status != 200:
            raise SystemExit(f"GET {url} -> HTTP {resp.status}")
        return resp.read()


def sha256(raw):
    return hashlib.sha256(raw).hexdigest()


def check_sha(name, raw, expected):
    got = sha256(raw)
    if got != expected:
        raise SystemExit(
            f"{name}: sha256 {got} != pinned {expected}. The live file changed "
            "since it was verified; re-verify before trusting the target.")


def download_public_files():
    """Write the five public files into data/ and return their bytes."""
    raw_zip = fetch(TRAIN_ZIP_URL)
    check_sha("module4_exercise_train.zip", raw_zip, TRAIN_ZIP_SHA256)
    files = {}
    with zipfile.ZipFile(io.BytesIO(raw_zip)) as zf:
        names = sorted(zf.namelist())
        expected = sorted(STORE_FILES[s] for s in TRAIN_STORES)
        if names != expected:
            raise SystemExit(f"train zip holds {names}, expected {expected}")
        for name in names:
            files[name] = zf.read(name)  # CRC-checked by zipfile
    files[STORE_FILES[TEST_STORE]] = fetch(TEST_CSV_URL)
    for name in PUBLIC_FILES:
        check_sha(name, files[name], PUBLIC_SHA256[name])
        with open(os.path.join(DATA, name), "wb") as f:
            f.write(files[name])
    return files


def fetch_api_prices():
    auth = json.loads(fetch(API_AUTH_URL))
    password = auth["data"]["password"]
    prices = json.loads(fetch(API_PRICES_URL.format(password=password)))
    if prices["status"] != "success":
        raise SystemExit(f"prices API: {prices['status']} {prices['message']}")
    return prices["data"]


def load_scrape_source():
    """The exercise table of /module4/scrapable-data, from the React source."""
    with open(SCRAPE_JS, encoding="utf-8") as f:
        text = f.read()
    m = re.search(r"export const scrapableDataExercise = (\[.*?\]);", text,
                  flags=re.S)
    if m is None:
        raise SystemExit(f"scrapableDataExercise not found in {SCRAPE_JS}")
    return json.loads(m.group(1))


def load_scrape_live():
    """The same records, from the live site bundle the page renders from."""
    html = fetch(SCRAPE_PAGE_URL).decode("utf-8")
    bundles = re.findall(r'src="(/static/js/main\.[0-9a-f]+\.js)"', html)
    if len(bundles) != 1:
        raise SystemExit(f"{SCRAPE_PAGE_URL}: expected one main bundle, "
                         f"found {bundles}")
    js = fetch(SITE + bundles[0]).decode("utf-8")
    rows = re.findall(r'\{item_code:"(P\d{4})",customer_score:(\d+),'
                      r'total_reviews:(\d+)\}', js)
    return [{"item_code": c, "customer_score": int(s),
             "total_reviews": int(r)} for c, s, r in rows]


# --------------------------------------------------------------------------- #
# The generator, module4_gen-exo.ipynb, replayed in its execution order
# --------------------------------------------------------------------------- #
def run_generator():
    # ---- cell 0 (verbatim apart from the final print / to_csv) ----
    np.random.seed(42)
    n_samples = 2000
    stores = ['CityMart', 'Greenfield_Grocers', 'SuperSaver_Outlet',
              'HighStreet_Bazaar', 'Neighborhood_Market']
    data = {
        'item_code': [f'P{i:04d}' for i in range(1, n_samples + 1)],
        'store_name': np.random.choice(stores, n_samples),
        'mass': np.round(np.random.uniform(0.1, 10.0, n_samples), 2),
        'dimension_length': np.round(np.random.uniform(5, 100, n_samples), 2),
        'dimension_width': np.round(np.random.uniform(5, 100, n_samples), 2),
        'dimension_height': np.round(np.random.uniform(5, 100, n_samples), 2),
        'customer_score': np.random.randint(1, 6, n_samples),
        'total_reviews': np.random.randint(0, 1000, n_samples),
        'days_since_last_purchase': np.random.randint(0, 365, n_samples),
    }
    df = pd.DataFrame(data)
    df['package_volume'] = (df['dimension_length'] * df['dimension_width']
                            * df['dimension_height'])
    df['stock_age'] = np.random.randint(1, 1000, n_samples)
    price = (15 + 0.5 * df['mass'] + np.random.normal(0, 5, n_samples))
    df['unit_cost'] = np.round(np.maximum(price, 1), 2)

    def sales_function(row, store_effect):
        base_sales = (
            150 +
            2 * row['mass'] +
            0.00007 * row['package_volume'] +
            10 * row['customer_score'] +
            0.05 * row['total_reviews'] +
            -0.05 * row['days_since_last_purchase'] +
            -2 * row['unit_cost']
        )
        return np.round(np.maximum(
            base_sales * store_effect + np.random.normal(0, 20), 0)
        ).astype(int)

    store_sales_effects = STORE_EFFECT
    df['quantity_sold'] = df.apply(
        lambda row: sales_function(row, store_sales_effects[row['store_name']]),
        axis=1)
    start_date = datetime(2023, 1, 1)
    df['last_modified'] = [start_date + timedelta(days=x)
                           for x in range(n_samples)]
    df_ad = df[df['store_name'].isin(TRAIN_STORES)]
    df_e = df[df['store_name'] == 'Neighborhood_Market']
    features = ['mass', 'dimension_length', 'dimension_width',
                'dimension_height', 'package_volume', 'customer_score',
                'total_reviews', 'days_since_last_purchase', 'unit_cost']
    model = LinearRegression()
    model.fit(df_ad[features], df_ad['quantity_sold'])
    raw_pred = model.predict(df_e[features])
    shift = np.random.normal(0, 6)  # ONE draw, added to every item
    df.loc[df['store_name'] == 'Neighborhood_Market', 'quantity_sold'] = (
        np.round(np.maximum(raw_pred.astype(int) + shift, 0)).astype(int))
    full = df.copy()

    # ---- cell 11 ----
    df = df[['item_code', 'store_name', 'mass', 'dimension_length',
             'dimension_width', 'dimension_height',
             'days_since_last_purchase', 'package_volume', 'stock_age',
             'quantity_sold', 'last_modified']]
    # ---- cell 14 ----
    store_dfs = {store: df[df['store_name'] == store] for store in stores}

    # ---- cell 15 ----
    def add_null(df, n=3, columns=['mass', 'dimension_length',
                                   'dimension_width',
                                   'days_since_last_purchase', 'stock_age',
                                   'package_volume']):
        selected_column = np.random.choice(columns)
        indices = np.random.choice(df.index, size=n, replace=False)
        df = df.copy()  # the notebook wrote into the slice; same values
        df.loc[indices, selected_column] = np.nan
        return df

    outputs, null_draws = {}, {}

    def nulled(store):
        before = store_dfs[store]
        after = add_null(before)
        col = [c for c in after.columns
               if after[c].isna().any() and not before[c].isna().any()]
        null_draws[store] = (col[0], sorted(after.loc[after[col[0]].isna(),
                                                      'item_code']))
        return after

    # ---- cell 16: CityMart ----
    store = 'CityMart'
    store_dfs[store] = nulled(store)
    outputs[store] = store_dfs[store].to_csv(index=False).encode()

    # ---- cell 17: Greenfield_Grocers ----
    store = 'Greenfield_Grocers'
    store_dfs[store] = nulled(store)
    store_dfs[store].columns = [col.upper() for col in store_dfs[store].columns]
    store_dfs[store]['1'] = ''
    store_dfs[store][''] = ''
    buf = io.StringIO()
    for _ in range(3):
        buf.write(f'{"|" * (len(store_dfs[store].columns) - 1)}\n')
    store_dfs[store].to_csv(buf, index=False, sep='|')
    outputs[store] = buf.getvalue().encode()

    # ---- cell 18: SuperSaver_Outlet (the grid the author's pandas wrote) ----
    store = 'SuperSaver_Outlet'
    store_dfs[store] = nulled(store)
    first_sheet_columns = ['item_code', 'quantity_sold']
    second_sheet_columns = ['item_code', 'store_name', 'mass',
                            'dimension_length', 'dimension_width',
                            'dimension_height', 'days_since_last_purchase',
                            'package_volume', 'stock_age']
    second_sheet_labels = ['item code', 'store name', 'mass',
                           'dimension length', 'dimension width',
                           'dimension height', 'days_since last_purchase',
                           'package volume', 'stock age']
    sheet = store_dfs[store]
    outputs[store] = {
        'Quantity': [first_sheet_columns]
        + sheet[first_sheet_columns].values.tolist(),
        'Info': [[None] + second_sheet_labels]
        + [row + [None] for row in sheet[second_sheet_columns].values.tolist()],
    }

    # ---- cell 19: HighStreet_Bazaar ----
    store = 'HighStreet_Bazaar'
    store_dfs[store] = nulled(store)
    # to_json's default when the file was written was epoch milliseconds for
    # datetimes; that path is deprecated, so the column is converted first.
    epoch_ms = ((store_dfs[store]['last_modified'] - pd.Timestamp(0))
                // pd.Timedelta(milliseconds=1))
    outputs[store] = store_dfs[store].assign(last_modified=epoch_ms).to_json(
        orient='records').encode()

    # ---- cell 20: Neighborhood_Market ----
    store = 'Neighborhood_Market'
    x = store_dfs[store].drop('quantity_sold', axis=1)
    outputs[store] = x.to_csv(index=False).encode()
    target = store_dfs[store][['item_code', 'quantity_sold']].reset_index(
        drop=True)

    return {"full": full, "target": target, "outputs": outputs,
            "null_draws": null_draws, "shift": shift,
            "store_dfs": store_dfs}


# --------------------------------------------------------------------------- #
# Checks
# --------------------------------------------------------------------------- #
def first_difference(a, b):
    la, lb = a.decode().splitlines(), b.decode().splitlines()
    for i, (x, y) in enumerate(zip(la, lb)):
        if x != y:
            return f"line {i + 1}:\n  download:    {x}\n  regenerated: {y}"
    return f"line counts differ: download {len(la)}, regenerated {len(lb)}"


def cell_equal(got, expected):
    """One xlsx cell. openpyxl stores numbers as "%.16g"
    (openpyxl.compat.strings.safe_string), so 30825.854399999997 is written
    30825.8544; the regenerated value goes through the same formatting."""
    if got is None or expected is None:
        return got is None and expected is None
    if isinstance(got, str) or isinstance(expected, str):
        return got == expected
    return float(got) == float("%.16g" % expected)


def check_store_files(gen, files):
    for store in [*TRAIN_STORES, TEST_STORE]:
        name = STORE_FILES[store]
        expected = gen["outputs"][store]
        if name.endswith(".xlsx"):
            wb = openpyxl.load_workbook(io.BytesIO(files[name]))
            if wb.sheetnames != list(expected):
                raise SystemExit(f"{name}: sheets {wb.sheetnames}, "
                                 f"expected {list(expected)}")
            for sheet, rows in expected.items():
                got = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
                if len(got) != len(rows):
                    raise SystemExit(f"{name}[{sheet}]: {len(got)} rows, "
                                     f"regenerated {len(rows)}")
                for i, (g, e) in enumerate(zip(got, rows)):
                    e = [None if isinstance(v, float) and np.isnan(v) else v
                         for v in e]
                    if len(g) != len(e) or not all(map(cell_equal, g, e)):
                        raise SystemExit(f"{name}[{sheet}] row {i + 1}:\n"
                                         f"  download:    {g}\n"
                                         f"  regenerated: {e}")
        elif files[name] != expected:
            raise SystemExit(f"{name} differs from the regenerated file at "
                             + first_difference(files[name], expected))
        print(f"  ok  {name} reproduced "
              f"{'cell by cell' if name.endswith('.xlsx') else 'byte for byte'}")


def check_side_sources(full, prices, scrape_source, scrape_live):
    codes = full["item_code"].tolist()
    if sorted(prices) != codes:
        raise SystemExit(f"prices API: {len(prices)} items, item codes differ "
                         "from P0001..P2000")
    api = pd.Series(prices).reindex(codes).to_numpy(dtype=float)
    bad = np.flatnonzero(api != full["unit_cost"].to_numpy())
    if bad.size:
        i = bad[0]
        raise SystemExit(f"unit_cost differs on {bad.size} items, first "
                         f"{codes[i]}: API {api[i]} vs generator "
                         f"{full['unit_cost'].iloc[i]}")
    print("  ok  unit_cost == live prices API, 2000/2000 items")

    expected = full[["item_code", "customer_score", "total_reviews"]]
    for label, records in [("ScrapableDataExercise.js", scrape_source),
                           ("live site bundle", scrape_live)]:
        got = pd.DataFrame(records)
        if len(got) != 2000:
            raise SystemExit(f"{label}: {len(got)} records, expected 2000")
        got = got.sort_values("item_code").reset_index(drop=True)
        diff = got.ne(expected.astype({"customer_score": "int64",
                                       "total_reviews": "int64"})).any(axis=1)
        if diff.any():
            raise SystemExit(f"{label}: {int(diff.sum())} records differ, "
                             f"first:\n{got[diff].head(3)}\n"
                             f"{expected[diff].head(3)}")
        print(f"  ok  customer_score, total_reviews == {label}, 2000/2000")


def check_target(gen):
    full, target = gen["full"], gen["target"]
    head = target.set_index("item_code")["quantity_sold"]
    for code, value in NOTEBOOK_TARGET_HEAD.items():
        if head[code] != value:
            raise SystemExit(f"target {code} = {head[code]}, the notebook "
                             f"printed {value}")
    # module4_gen-exo.ipynb cell 6, last line.
    dfe = full[full["store_name"] == TEST_STORE]
    train = full[full["store_name"] != TEST_STORE]
    lr = LinearRegression()
    lr.fit(train.iloc[:, 2:-2], train.iloc[:, -2])
    mae = mean_absolute_error(lr.predict(dfe.iloc[:, 2:-2]), dfe.iloc[:, -2])
    if abs(mae - NOTEBOOK_FINAL_TEST_MAE) > 1e-6:
        raise SystemExit(f"cell 6 'final test on E' = {mae!r}, the notebook "
                         f"printed {NOTEBOOK_FINAL_TEST_MAE!r}")
    print(f"  ok  target head and cell 6 'final test on E' = {mae:.12f} "
          f"(notebook {NOTEBOOK_FINAL_TEST_MAE})")


# --------------------------------------------------------------------------- #
# The ladder — module4_exercise1.ipynb's get_simple_baseline, verbatim
# --------------------------------------------------------------------------- #
def get_simple_baseline(data, fillna_value=-1, drop_cols=None, k_fold=5,
                        scaler='standard', model='linear', metric='mae',
                        target_col=None, X_data_test=None):
    data = data.copy()
    data.fillna(fillna_value, inplace=True)
    if X_data_test is not None:
        X_data_test = X_data_test.copy()
        X_data_test.fillna(fillna_value, inplace=True)
    if drop_cols:
        data.drop(drop_cols, axis=1, inplace=True)
        if X_data_test is not None:
            X_data_test.drop(drop_cols, axis=1, inplace=True)
    y = data[target_col]
    X = data.drop(target_col, axis=1)
    if scaler != 'standard':
        raise ValueError("only the standard scaler is used here")
    scaler = StandardScaler()
    X = scaler.fit_transform(X)
    if X_data_test is not None:
        X_data_test = scaler.transform(X_data_test)
    if model != 'linear' or metric != 'mae':
        raise ValueError("only linear / mae are used here")
    model = LinearRegression()
    kf = KFold(n_splits=k_fold, shuffle=True, random_state=42)
    scores = []
    for train_index, test_index in kf.split(X):
        model.fit(X[train_index], y.iloc[train_index])
        scores.append(mean_absolute_error(y.iloc[test_index],
                                          model.predict(X[test_index])))
    if X_data_test is not None:
        model.fit(X, y)
        return np.mean(scores), model.predict(X_data_test)
    return np.mean(scores)


def assemble_like_a_student(prices, scrape_source):
    """The public sources, read and merged as the reference solution did."""
    a = pd.read_csv(os.path.join(DATA, "CityMart_data.csv"),
                    index_col="item_code")
    b = pd.read_csv(os.path.join(DATA, "Greenfield_Grocers_data.csv"),
                    sep="|", header=3, index_col="ITEM_CODE")
    b = b.drop(["1", "Unnamed: 12"], axis=1)
    b.columns = [c.lower() for c in b.columns]
    b.index.name = "item_code"
    cs = pd.read_excel(os.path.join(DATA, "SuperSaver_Outlet_data.xlsx"),
                       sheet_name=None)
    info = cs["Info"]
    info.columns = ["item_code", "store_name", *FILE_COLS, "toremove"]
    info = info.drop("toremove", axis=1)
    c = pd.merge(cs["Quantity"], info, on="item_code").set_index("item_code")
    d = pd.read_json(os.path.join(DATA, "HighStreet_Bazaar_data.json"),
                     orient="records").set_index("item_code")
    train = pd.concat([a, b, c, d], axis=0)
    test = pd.read_csv(os.path.join(DATA, "Neighborhood_Market_data.csv"),
                       index_col="item_code")
    api = pd.DataFrame.from_dict(prices, orient="index", columns=["unit_cost"])
    scraped = pd.DataFrame(scrape_source).set_index("item_code")
    scraped = scraped.astype(float)  # the reference scraper parsed float()
    train = train.join(api, how="left").join(scraped, how="left")
    test = test.join(api, how="left").join(scraped, how="left")
    for frame in (train, test):
        frame["weekly_footfall"] = frame["store_name"].map(WEEKLY_FOOTFALL)
    return train, test


def run_ladder(train, test, target):
    y_true = target.set_index("item_code")["quantity_sold"].reindex(test.index)
    if y_true.isna().any():
        raise SystemExit("test items without a target")
    public_board = np.array([int(code[1:]) % 2 == 0 for code in test.index])
    drop = ["store_name", "last_modified"]
    rungs = [
        ("files", FILE_COLS),
        ("files+api", FILE_COLS + ["unit_cost"]),
        ("files+api+scraping",
         FILE_COLS + ["unit_cost", "customer_score", "total_reviews"]),
        ("files+api+scraping+footfall",
         FILE_COLS + ["unit_cost", "customer_score", "total_reviews",
                      "weekly_footfall"]),
    ]
    ladder = {}

    def record(name, pred, cv=None):
        err = np.abs(np.asarray(pred, dtype=float) - y_true.to_numpy())
        ladder[name] = {"mae": float(err.mean()),
                        "mae_public_board": float(err[public_board].mean()),
                        "cv_mae": None if cv is None else float(cv),
                        "passes": bool(err.mean() <= PASS_MAE)}

    record("constant_train_mean",
           np.full(len(test), train["quantity_sold"].mean()))
    for name, cols in rungs:
        cv, pred = get_simple_baseline(
            train[["store_name", "last_modified", "quantity_sold", *cols]],
            drop_cols=drop, target_col="quantity_sold",
            X_data_test=test[["store_name", "last_modified", *cols]])
        if name in REFERENCE_CV_MAE and abs(cv - REFERENCE_CV_MAE[name]) > 1e-6:
            raise SystemExit(f"CV MAE for {name} = {cv!r}, the reference "
                             f"solution printed {REFERENCE_CV_MAE[name]!r}")
        record(name, pred, cv)

    # Footfall as a multiplier: learn sales per visitor, rescale per store.
    cols = rungs[2][1]
    per_visit = train[["store_name", "last_modified", *cols]].copy()
    per_visit["quantity_sold"] = (train["quantity_sold"]
                                  / train["weekly_footfall"])
    _, pred = get_simple_baseline(
        per_visit, drop_cols=drop, target_col="quantity_sold",
        X_data_test=test[["store_name", "last_modified", *cols]])
    record("footfall_multiplier", pred * test["weekly_footfall"].to_numpy())
    return ladder


# --------------------------------------------------------------------------- #
def nan_pattern(files):
    out = {}
    readers = {
        "CityMart_data.csv": lambda r: pd.read_csv(io.BytesIO(r)),
        "Greenfield_Grocers_data.csv": lambda r: pd.read_csv(
            io.BytesIO(r), sep="|", header=3).drop(
                columns=["1", "Unnamed: 12"]),
        "HighStreet_Bazaar_data.json": lambda r: pd.read_json(
            io.BytesIO(r), orient="records"),
        "Neighborhood_Market_data.csv": lambda r: pd.read_csv(io.BytesIO(r)),
    }
    for name, read in readers.items():
        frame = read(files[name])
        key = frame.columns[0]
        out[name] = {col: frame.loc[frame[col].isna(), key].tolist()
                     for col in frame.columns if frame[col].isna().any()}
    cs = pd.read_excel(io.BytesIO(files["SuperSaver_Outlet_data.xlsx"]),
                       sheet_name=None, header=None)
    info = cs["Info"].iloc[1:, :9]  # the data cells, left of the shift
    info.columns = ["item_code", "store_name", *FILE_COLS]
    out["SuperSaver_Outlet_data.xlsx"] = {
        f"Info/{col}": info.loc[info[col].isna(), "item_code"].tolist()
        for col in info.columns if info[col].isna().any()}
    return out


def main():
    os.makedirs(DATA, exist_ok=True)

    print("downloading the public files …")
    files = download_public_files()
    for name in PUBLIC_FILES:
        print(f"  {name:32s} {len(files[name]):7d} bytes  sha256 ok")

    print("fetching the prices API and the scraped table …")
    prices = fetch_api_prices()
    scrape_source = load_scrape_source()
    scrape_live = load_scrape_live()

    print("replaying module4_gen-exo.ipynb …")
    with warnings.catch_warnings():
        warnings.simplefilter("error")  # a pandas behaviour change must stop us
        gen = run_generator()
    for store, (col, codes) in gen["null_draws"].items():
        print(f"  add_null {store:20s} {col:26s} {codes}")

    print("checking …")
    check_store_files(gen, files)
    check_side_sources(gen["full"], prices, scrape_source, scrape_live)
    check_target(gen)

    full = gen["full"].copy()
    full["last_modified"] = full["last_modified"].dt.strftime("%Y-%m-%d")
    full["weekly_footfall"] = full["store_name"].map(WEEKLY_FOOTFALL)
    target = gen["target"]

    train, test = assemble_like_a_student(prices, scrape_source)
    ladder = run_ladder(train, test, target)
    print("  ok  reference solution's three CV scores reproduced")

    target.to_csv(os.path.join(DATA, "neighborhood_market_target.csv"),
                  index=False)
    shutil.copyfile(os.path.join(DATA, "neighborhood_market_target.csv"),
                    os.path.join(DATA, "y_test.csv"))
    full.to_csv(os.path.join(DATA, "full_generated.csv"), index=False)
    with open(os.path.join(DATA, "ladder.json"), "w") as f:
        json.dump(ladder, f, indent=2)

    print("\nrows per store")
    for store, n in full["store_name"].value_counts().sort_index().items():
        print(f"  {store:20s} {n}")
    q = target["quantity_sold"]
    print(f"\ntarget: {len(q)} items, mean {q.mean():.4f}, std {q.std():.4f}, "
          f"min {q.min()}, max {q.max()}; single shift draw "
          f"{gen['shift']:.6f}")
    print("\nNaN pattern per public file")
    for name, pattern in nan_pattern(files).items():
        print(f"  {name}: {pattern if pattern else 'none'}")
    print(f"\nladder (test MAE on {len(test)} items; public board = even "
          f"item numbers; pass <= {PASS_MAE})")
    for name, row in ladder.items():
        cv = "" if row["cv_mae"] is None else f"  cv {row['cv_mae']:.4f}"
        print(f"  {name:30s} {row['mae']:9.4f}  public "
              f"{row['mae_public_board']:9.4f}  "
              f"{'pass' if row['passes'] else 'fail'}{cv}")
    print(f"\nwrote {DATA}")
    print("  public : " + " ".join(PUBLIC_FILES))
    print("  private: neighborhood_market_target.csv y_test.csv (identical) "
          "full_generated.csv ladder.json")


if __name__ == "__main__":
    main()
