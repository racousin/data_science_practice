#!/usr/bin/env python3
"""Build PPTX decks from the Markdown courseware.

One deck per module (= one 3h session). Reads ``course.yaml``, walks the
modules in manifest order, and renders every lesson body into slides. A lesson
with ``in_deck: false`` is skipped — self-study reference material that lives in
the module for the web course but is never lectured from it.

Slide model
-----------
A lesson body is split on ``---`` lines (a Markdown horizontal rule, which the
ML-Arena renderer shows as a separator — so the *same* file is a valid web
lesson and a valid deck). Inside a slide:

* the first ATX heading becomes the slide title,
* everything else becomes the body, laid out top-down,
* ``<!-- notes: ... -->`` becomes the speaker-notes pane.

Usage::

    uv run --with python-pptx --with pyyaml \\
        python tools/build_slides.py content/python-ai-engineering --out build/slides
"""

from __future__ import annotations

import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import yaml
from mathrender import MathRenderer, UnknownMacro, latex_to_unicode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# --------------------------------------------------------------------------- #
# Theme
# --------------------------------------------------------------------------- #

INK = RGBColor(0x1A, 0x1B, 0x1E)
MUTED = RGBColor(0x5C, 0x63, 0x70)
ACCENT = RGBColor(0x1C, 0x6E, 0xB8)
RULE = RGBColor(0xDE, 0xE2, 0xE6)
CODE_BG = RGBColor(0x1E, 0x22, 0x2A)
CODE_FG = RGBColor(0xE6, 0xEA, 0xF0)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)

BODY_FONT = "Helvetica Neue"
MONO_FONT = "Menlo"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.75)
CONTENT_W = SLIDE_W - 2 * MARGIN
CONTENT_W_IN = CONTENT_W / Emu(914400)
BODY_TOP = Inches(1.85)
BODY_BOTTOM = SLIDE_H - Inches(0.7)

# Body font sizes, largest first. The renderer measures the slide and steps
# down this ladder until the content fits.
SIZE_LADDER = [20, 18, 16, 14, 12, 11, 10]

# A figure is the one block on a slide that can give ground: it takes whatever
# height the text leaves it and no more. Alone on a slide it grows to fill the
# body; sharing with text it shrinks, keeping IMAGE_GOOD_H for as long as a rung
# of the ladder allows and bottoming out at IMAGE_MIN_H — below which it stops
# being readable from the back of a room and the slide has to be split instead.
IMAGE_MIN_H = 2.2   # inches; the floor, past which check_slide_overflow speaks up
IMAGE_GOOD_H = 3.0  # inches; kept in preference to a larger body font
IMAGE_BASE_H = 3.4  # inches; every figure may reach this regardless of resolution
IMAGE_MAX_H = 4.6   # inches; a figure alone on a slide fills the body to here
IMAGE_MIN_DPI = 110  # do not enlarge a small figure past this


# --------------------------------------------------------------------------- #
# Markdown block parsing
# --------------------------------------------------------------------------- #

NOTES_RE = re.compile(r"<!--\s*notes:\s*(.*?)-->", re.S | re.I)
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
ORDERED_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
FENCE_RE = re.compile(r"^```([\w+-]*)\s*$")
TABLE_SEP_RE = re.compile(r"^\|[\s:|-]+\|$")


class Block:
    """One renderable unit inside a slide."""

    def __init__(self, kind: str, **kw):
        self.kind = kind
        self.__dict__.update(kw)

    def __repr__(self) -> str:  # pragma: no cover - debugging aid
        return f"Block({self.kind})"


def split_slides(body: str) -> list[str]:
    """Split a lesson body on horizontal rules, ignoring rules inside fences."""
    slides, current, in_fence = [], [], False
    for line in body.splitlines():
        if FENCE_RE.match(line.strip()) or line.strip() == "```":
            in_fence = not in_fence
        if not in_fence and re.fullmatch(r"-{3,}|\*{3,}|_{3,}", line.strip()):
            slides.append("\n".join(current))
            current = []
            continue
        current.append(line)
    slides.append("\n".join(current))
    return [s for s in slides if s.strip()]


def parse_blocks(text: str) -> tuple[str | None, list[Block]]:
    """Parse one slide's Markdown into (title, blocks)."""
    lines = text.splitlines()
    title: str | None = None
    blocks: list[Block] = []
    para: list[str] = []
    i = 0

    def flush_para() -> None:
        nonlocal para
        joined = " ".join(l.strip() for l in para).strip()
        if joined:
            blocks.append(Block("para", text=joined))
        para = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush_para()
            i += 1
            continue

        fence = FENCE_RE.match(stripped)
        if fence:
            flush_para()
            lang = fence.group(1)
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(Block("code", lang=lang, lines=_dedent(code)))
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            flush_para()
            level, text_ = len(heading.group(1)), heading.group(2).strip()
            if title is None:
                title = text_
            else:
                blocks.append(Block("subhead", text=text_, level=level))
            i += 1
            continue

        img = IMAGE_RE.match(stripped)
        if img:
            flush_para()
            blocks.append(Block("image", alt=img.group(1), src=img.group(2)))
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            flush_para()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = lines[i].strip()
                if not TABLE_SEP_RE.match(row):
                    rows.append([c.strip() for c in row.strip("|").split("|")])
                i += 1
            if rows:
                blocks.append(Block("table", rows=rows))
            continue

        if stripped.startswith("$$"):
            flush_para()
            i += 1
            math: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("$$"):
                math.append(lines[i].strip())
                i += 1
            i += 1
            blocks.append(Block("math", text=" ".join(math)))
            continue

        if stripped.startswith(">"):
            flush_para()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(Block("quote", text=" ".join(quote)))
            continue

        bullet = BULLET_RE.match(line)
        ordered = ORDERED_RE.match(line)
        if bullet or ordered:
            flush_para()
            items: list[tuple[int, str, str | None]] = []
            while i < len(lines):
                b = BULLET_RE.match(lines[i])
                o = ORDERED_RE.match(lines[i])
                if b:
                    items.append((len(b.group(1)) // 2, b.group(2).strip(), None))
                elif o:
                    items.append((len(o.group(1)) // 2, o.group(3).strip(), o.group(2)))
                elif lines[i].strip() and lines[i].startswith(("  ", "\t")) and items:
                    lvl, txt, marker = items[-1]
                    items[-1] = (lvl, f"{txt} {lines[i].strip()}", marker)
                else:
                    break
                i += 1
            blocks.append(Block("list", items=items))
            continue

        para.append(line)
        i += 1

    flush_para()
    return title, blocks


def _dedent(lines: list[str]) -> list[str]:
    body = [l for l in lines if l.strip()]
    if not body:
        return lines
    pad = min(len(l) - len(l.lstrip()) for l in body)
    return [l[pad:] if len(l) >= pad else l for l in lines]


# --------------------------------------------------------------------------- #
# Inline markdown -> pptx runs
# --------------------------------------------------------------------------- #

INLINE_RE = re.compile(
    r"(?P<code>`[^`]+`)"
    r"|(?P<bold>\*\*[^*]+\*\*)"
    r"|(?P<italic>(?<!\*)\*[^*]+\*(?!\*))"
    r"|(?P<link>\[[^\]]+\]\([^)]+\))"
    r"|(?P<math>\$[^$]+\$)"
)


def write_inline(paragraph, text: str, size: int, color=INK, bold=False,
                 italic=False) -> None:
    """Render inline Markdown into runs on an existing pptx paragraph.

    Emphasis spans recurse so that a link inside `*...*` (or code inside
    `**...**`) is still parsed rather than printed with its markers.
    """
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            _run(paragraph, text[pos : m.start()], size, color, bold, italic)
        kind = m.lastgroup
        raw = m.group()
        if kind == "code":
            _run(paragraph, raw[1:-1], size - 1, ACCENT, bold, italic, mono=True)
        elif kind == "bold":
            write_inline(paragraph, raw[2:-2], size, color, True, italic)
        elif kind == "italic":
            write_inline(paragraph, raw[1:-1], size, color, bold, True)
        elif kind == "link":
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", raw).groups()
            run = _run(paragraph, label, size, ACCENT, bold, italic)
            try:
                run.hyperlink.address = url
            except Exception:
                pass
        elif kind == "math":
            _run(paragraph, latex_to_unicode(raw[1:-1]), size, color, bold,
                 italic=True)
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], size, color, bold, italic)


def _run(paragraph, text: str, size: int, color, bold: bool, italic=False, mono=False):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = MONO_FONT if mono else BODY_FONT
    run.font.color.rgb = color
    return run


# --------------------------------------------------------------------------- #
# Height estimation (drives the auto-fit ladder)
# --------------------------------------------------------------------------- #

# Measured off Helvetica Neue: mixed-case prose averages 15.8 characters to the
# inch at 10pt (bold, 14.8). The model is set a little under that because a line
# breaks at a word boundary and so wastes up to one word of the last inch.
CHARS_PER_INCH = 14.0

# Matches the left/right cell margins _table sets.
TABLE_CELL_PAD_IN = 0.08


def _wrapped_lines(text: str, size: int, width_in: float) -> int:
    per_line = max(10, int(width_in * CHARS_PER_INCH * (10.0 / size)))
    return max(1, -(-len(_plain(text)) // per_line))


def _plain(text: str) -> str:
    return re.sub(r"[*`$]|\[|\]\([^)]*\)", "", text)


class Measurer:
    """How tall each block renders, in inches.

    The renderer and ``check_slide_overflow`` share one of these, so the height
    the auto-fit ladder searches over is the height the slide actually gets. The
    three things a flat per-block guess got wrong are all resolved from the real
    thing here: an image from its file, display math from the PNG mathtext
    produces, and a table row from the text that has to wrap inside it.
    """

    def __init__(self, asset_roots: list[str], math: MathRenderer | None = None):
        self.asset_roots = list(asset_roots)
        self.math = math
        self._pixels: dict[str, tuple[int, int] | None] = {}

    # -- assets ------------------------------------------------------------ #
    def resolve(self, src: str) -> str | None:
        if src.startswith(("http://", "https://")):
            return None
        for root in self.asset_roots:
            candidate = os.path.normpath(os.path.join(root, src.lstrip("/")))
            if os.path.isfile(candidate):
                return candidate
        return None

    def pixels(self, src: str) -> tuple[int, int] | None:
        """An image's pixel size, or None if it cannot be resolved."""
        if src not in self._pixels:
            path = self.resolve(src)
            size = None
            if path is not None:
                from PIL import Image as PILImage  # ships with python-pptx

                with PILImage.open(path) as im:
                    size = im.size
            self._pixels[src] = size
        return self._pixels[src]

    def image_cap(self, block: Block) -> float:
        """The tallest this figure may be drawn.

        Bounded three ways: it may not be wider than the body box, it may not be
        enlarged past IMAGE_MIN_DPI (a 300 px screenshot blown up to fill a slide
        looks worse than a small sharp one), and it may not exceed IMAGE_MAX_H.
        IMAGE_BASE_H is the floor on the resolution rule, so no figure is ever
        drawn smaller than it used to be.
        """
        size = self.pixels(block.src)
        if size is None:
            return 0.60  # renders as a one-line "[missing image]" placeholder
        px_w, px_h = size
        return min(
            IMAGE_MAX_H,
            CONTENT_W_IN / (px_w / px_h),
            max(IMAGE_BASE_H, px_h / IMAGE_MIN_DPI),
        )

    def image_floor(self, block: Block) -> float:
        return min(IMAGE_MIN_H, self.image_cap(block))

    def image_comfort(self, block: Block) -> float:
        return min(IMAGE_GOOD_H, self.image_cap(block))

    # -- blocks ------------------------------------------------------------ #
    def math_height(self, block: Block, size: int) -> float:
        """The rendered PNG's height, scaled the way _math scales it."""
        if self.math is None:
            return 0.62
        from PIL import Image as PILImage

        with PILImage.open(self.math.render(block.text, fontsize=size + 6)) as im:
            px_w, px_h = im.size
            dpi_x, dpi_y = im.info.get("dpi", (72.0, 72.0))
        w_in, h_in = px_w / dpi_x, px_h / dpi_y
        return h_in * min(1.0, (CONTENT_W_IN * 0.75) / w_in)

    def table_row_heights(self, rows: list[list[str]], size: int) -> list[float]:
        """One height per row, tall enough for the row's most wrapped cell.

        PowerPoint grows a row that its text does not fit in, which used to push
        everything below the table down the slide without anyone measuring it.
        """
        cols = max(len(r) for r in rows)
        fs = max(9, size - 3)
        col_in = CONTENT_W_IN / cols - 2 * TABLE_CELL_PAD_IN
        line_in = fs * 1.30 / 72.0
        return [
            max(
                size * 1.9 / 72.0,
                max(_wrapped_lines(c, fs, col_in) for c in row) * line_in + 0.10,
            )
            for row in rows
        ]

    def block_height(self, block: Block, size: int) -> float:
        """Height of one non-image block, spacing included.

        Every arm mirrors what the matching renderer method advances ``y`` by.
        Images are absent on purpose — they are allocated together, in `plan`.
        """
        line_h = size * 1.45 / 72.0
        b = block
        if b.kind == "para":
            return _wrapped_lines(b.text, size, CONTENT_W_IN) * line_h + 0.10
        if b.kind == "subhead":
            return (size + 3) * 1.5 / 72.0 + 0.12
        if b.kind == "list":
            return sum(
                _wrapped_lines(t, size, CONTENT_W_IN - lvl * 0.35) * line_h + 0.05
                for lvl, t, _ in b.items
            ) + 0.10
        if b.kind == "code":
            return len(b.lines) * max(9, size - 2) * 1.35 / 72.0 + 0.28 + 0.14
        if b.kind == "table":
            return sum(self.table_row_heights(b.rows, size)) + 0.16
        if b.kind == "math":
            return self.math_height(b, size) + 0.20
        if b.kind == "quote":
            return _wrapped_lines(b.text, size, CONTENT_W_IN - 0.4) * line_h + 0.32
        raise ValueError(f"unmeasurable block kind: {b.kind!r}")

    # -- whole slide -------------------------------------------------------- #
    def plan(self, blocks: list[Block], size: int, avail: float
             ) -> tuple[float, dict[int, float]]:
        """Lay a slide out at `size`: (total height, height per image block).

        Text is incompressible. Figures are not, so they are handed what the
        text leaves, shared equally, capped at their natural height and floored
        at IMAGE_MIN_H — below which shrinking stops and the total runs past
        `avail`, which is exactly what "this slide overflows" means.
        """
        images = [(i, b) for i, b in enumerate(blocks) if b.kind == "image"]
        fixed = sum(self.block_height(b, size) for b in blocks if b.kind != "image")
        if not images:
            return fixed, {}

        gaps = 0.16 * len(images)
        share = max(0.0, avail - fixed - gaps) / len(images)
        heights = [
            max(self.image_floor(b), min(self.image_cap(b), share))
            for _, b in images
        ]
        return (
            fixed + gaps + sum(heights),
            {i: h for (i, _), h in zip(images, heights)},
        )

    def fit(self, blocks: list[Block], avail: float
            ) -> tuple[int, dict[int, float], float]:
        """Pick the body size: (size, height per image block, total height).

        Largest rung that fits — except that a figure is worth more than two
        points of body text, so a rung only wins while it still leaves every
        figure its comfortable height. When no rung does, the largest one that
        merely fits is taken, and when none of them fit the smallest is used and
        `check_slide_overflow` reports the slide.
        """
        fitting = None
        for size in SIZE_LADDER:
            total, heights = self.plan(blocks, size, avail)
            if total > avail:
                continue
            if fitting is None:
                fitting = (size, heights, total)
            if all(h >= self.image_comfort(blocks[i]) for i, h in heights.items()):
                return size, heights, total
        if fitting is not None:
            return fitting
        size = SIZE_LADDER[-1]
        total, heights = self.plan(blocks, size, avail)
        return size, heights, total


# --------------------------------------------------------------------------- #
# Deck rendering
# --------------------------------------------------------------------------- #


class DeckBuilder:
    def __init__(self, asset_roots: list[str], strict_assets: bool,
                 math: MathRenderer):
        self.math = math
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self.strict_assets = strict_assets
        self.missing_assets: list[str] = []
        self.measure = Measurer(asset_roots, math)

    # -- primitives -------------------------------------------------------- #
    def _blank(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        bg.fill.solid()
        bg.fill.fore_color.rgb = PAPER
        bg.line.fill.background()
        bg.shadow.inherit = False
        return slide

    def _textbox(self, slide, left, top, width, height):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        return tf

    # -- slide kinds -------------------------------------------------------- #
    def cover(self, title: str, subtitle: str, footer: str) -> None:
        slide = self._blank()
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.28), SLIDE_H)
        band.fill.solid()
        band.fill.fore_color.rgb = ACCENT
        band.line.fill.background()
        band.shadow.inherit = False

        tf = self._textbox(slide, Inches(1.1), Inches(2.4), SLIDE_W - Inches(2.2), Inches(2.6))
        p = tf.paragraphs[0]
        _run(p, title, 40, INK, True)
        if subtitle:
            p2 = tf.add_paragraph()
            p2.space_before = Pt(14)
            write_inline(p2, subtitle, 20, MUTED)
        if footer:
            p3 = tf.add_paragraph()
            p3.space_before = Pt(28)
            _run(p3, footer, 14, MUTED, False)

    def section(self, title: str, eyebrow: str) -> None:
        slide = self._blank()
        tf = self._textbox(slide, MARGIN, Inches(3.0), CONTENT_W, Inches(1.8))
        if eyebrow:
            p0 = tf.paragraphs[0]
            _run(p0, eyebrow.upper(), 13, ACCENT, True)
            p = tf.add_paragraph()
            p.space_before = Pt(8)
        else:
            p = tf.paragraphs[0]
        _run(p, title, 32, INK, True)

    def content(self, title: str | None, blocks: list[Block], notes: str, eyebrow: str) -> None:
        slide = self._blank()

        if title:
            tf = self._textbox(slide, MARGIN, Inches(0.62), CONTENT_W, Inches(0.9))
            p = tf.paragraphs[0]
            write_inline(p, title, 26, INK, bold=True)
            rule = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.55), Inches(1.6), Emu(22860)
            )
            rule.fill.solid()
            rule.fill.fore_color.rgb = ACCENT
            rule.line.fill.background()
            rule.shadow.inherit = False
            top = BODY_TOP
        else:
            top = Inches(1.0)

        if eyebrow and eyebrow.strip().lower() != (title or "").strip().lower():
            tf = self._textbox(slide, MARGIN, Inches(0.28), CONTENT_W, Inches(0.3))
            _run(tf.paragraphs[0], eyebrow, 10, MUTED, False)

        avail = (BODY_BOTTOM - top) / Emu(914400)
        size, image_heights, _ = self.measure.fit(blocks, avail)
        self._render_blocks(slide, blocks, top, size, image_heights)

        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    # -- block layout ------------------------------------------------------- #
    def _render_blocks(self, slide, blocks: list[Block], top, size: int,
                       image_heights: dict[int, float]) -> None:
        y = top
        for idx, b in enumerate(blocks):
            if b.kind == "para":
                y = self._para(slide, b.text, y, size)
            elif b.kind == "subhead":
                y = self._subhead(slide, b.text, y, size)
            elif b.kind == "list":
                y = self._list(slide, b.items, y, size)
            elif b.kind == "code":
                y = self._code(slide, b.lines, y, size)
            elif b.kind == "table":
                y = self._table(slide, b.rows, y, size)
            elif b.kind == "math":
                y = self._math(slide, b.text, y, size)
            elif b.kind == "quote":
                y = self._quote(slide, b.text, y, size)
            elif b.kind == "image":
                y = self._image(slide, b, y, image_heights[idx])

    def _para(self, slide, text: str, y, size: int):
        h = Inches(_wrapped_lines(text, size, CONTENT_W / Emu(914400)) * size * 1.45 / 72.0)
        tf = self._textbox(slide, MARGIN, y, CONTENT_W, h)
        write_inline(tf.paragraphs[0], text, size)
        return y + h + Inches(0.10)

    def _subhead(self, slide, text: str, y, size: int):
        h = Inches((size + 3) * 1.5 / 72.0)
        tf = self._textbox(slide, MARGIN, y + Inches(0.06), CONTENT_W, h)
        write_inline(tf.paragraphs[0], text, size + 3, ACCENT, bold=True)
        return y + h + Inches(0.12)

    def _list(self, slide, items, y, size: int):
        width_in = CONTENT_W / Emu(914400)
        total = sum(
            _wrapped_lines(t, size, width_in - lvl * 0.35) * size * 1.45 / 72.0 + 0.05
            for lvl, t, _ in items
        )
        tf = self._textbox(slide, MARGIN, y, CONTENT_W, Inches(total))
        for idx, (lvl, text, marker) in enumerate(items):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.space_after = Pt(4)
            indent = Inches(0.30 * lvl)
            p.left_indent = indent + Inches(0.24)
            p.first_line_indent = -Inches(0.24)
            bullet = f"{marker}. " if marker else ("•  " if lvl == 0 else "–  ")
            _run(p, bullet, size, ACCENT if lvl == 0 else MUTED, False)
            write_inline(p, text, size)
        return y + Inches(total) + Inches(0.10)

    def _code(self, slide, lines: list[str], y, size: int):
        csize = max(9, size - 2)
        h = Inches(len(lines) * csize * 1.35 / 72.0 + 0.28)
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, y, CONTENT_W, h)
        panel.fill.solid()
        panel.fill.fore_color.rgb = CODE_BG
        panel.line.fill.background()
        panel.shadow.inherit = False
        panel.adjustments[0] = 0.04

        tf = panel.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = Inches(0.18)
        tf.margin_top = tf.margin_bottom = Inches(0.12)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for idx, line in enumerate(lines):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(0)
            run = p.add_run()
            run.text = line or " "
            run.font.size = Pt(csize)
            run.font.name = MONO_FONT
            run.font.color.rgb = CODE_FG
        return y + h + Inches(0.14)

    def _table(self, slide, rows, y, size: int):
        cols = max(len(r) for r in rows)
        rows = [r + [""] * (cols - len(r)) for r in rows]
        row_heights = self.measure.table_row_heights(rows, size)
        h = Inches(sum(row_heights))
        shape = slide.shapes.add_table(len(rows), cols, MARGIN, y, CONTENT_W, h)
        table = shape.table
        # Set every row explicitly: add_table splits the total evenly, and a row
        # whose text needs two lines then grows past what was measured.
        for ri, rh in enumerate(row_heights):
            table.rows[ri].height = Inches(rh)
        for ri, row in enumerate(rows):
            for ci, cell_text in enumerate(row):
                cell = table.cell(ri, ci)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                tf = cell.text_frame
                tf.word_wrap = True
                write_inline(
                    tf.paragraphs[0], cell_text, max(9, size - 3),
                    INK if ri else PAPER, bold=(ri == 0),
                )
        return y + h + Inches(0.16)

    def _math(self, slide, text: str, y, size: int):
        """Display math goes in as a rendered picture — mathtext gives real
        fractions and summation signs, which no run of text can."""
        png = self.math.render(text, fontsize=size + 6)
        pic = slide.shapes.add_picture(png, MARGIN, y)
        scale = min(1.0, (CONTENT_W * 0.75) / pic.width)
        pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
        pic.left = int((SLIDE_W - pic.width) / 2)
        return y + pic.height + Inches(0.20)

    def _quote(self, slide, text: str, y, size: int):
        lines = _wrapped_lines(text, size, CONTENT_W / Emu(914400) - 0.4)
        h = Inches(lines * size * 1.45 / 72.0 + 0.16)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, y, Emu(34290), h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = RULE
        bar.line.fill.background()
        bar.shadow.inherit = False
        tf = self._textbox(slide, MARGIN + Inches(0.24), y + Inches(0.08),
                           CONTENT_W - Inches(0.24), h)
        write_inline(tf.paragraphs[0], text, size, MUTED)
        return y + h + Inches(0.16)

    def _image(self, slide, block: Block, y, height_in: float):
        path = self.measure.resolve(block.src)
        if path is None:
            self.missing_assets.append(block.src)
            return self._quote(slide, f"[missing image: {block.src}]", y, 14)
        h = min(Inches(height_in), BODY_BOTTOM - y)
        pic = slide.shapes.add_picture(path, MARGIN, y, height=h)
        if pic.width > CONTENT_W:  # guard; the plan already caps on width
            ratio = CONTENT_W / pic.width
            pic.width = int(pic.width * ratio)
            pic.height = int(pic.height * ratio)
        pic.left = int((SLIDE_W - pic.width) / 2)
        return y + pic.height + Inches(0.16)

    def save(self, path: str) -> None:
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
        self.prs.save(path)


# --------------------------------------------------------------------------- #
# Driver
# --------------------------------------------------------------------------- #


def load_manifest(base: str) -> dict:
    for name in ("course.yaml", "course.yml", "course.json"):
        path = os.path.join(base, name)
        if os.path.isfile(path):
            with open(path, encoding="utf-8") as fh:
                if name.endswith(".json"):
                    import json

                    return json.load(fh)
                return yaml.safe_load(fh)
    raise SystemExit(f"no course.yaml in {base}")


def build_module(base: str, course: dict, module: dict, out_dir: str,
                 strict_assets: bool, math: MathRenderer) -> tuple[str, int]:
    asset_roots = [base, os.path.join(base, "assets"), os.path.dirname(base)]
    deck = DeckBuilder(asset_roots, strict_assets, math)

    eyebrow = course.get("name", "")
    deck.cover(module["title"], module.get("summary", ""),
               f"{eyebrow} · {course.get('instructor_name', '')}".strip(" ·"))

    for lesson in module.get("lessons", []):
        if not lesson.get("in_deck", True):
            continue
        body = read_lesson_body(base, lesson)
        slides = split_slides(body)
        if not slides:
            continue
        deck.section(lesson["title"], lesson.get("kind", "lesson"))
        for n, raw in enumerate(slides, 1):
            notes_match = NOTES_RE.search(raw)
            notes = notes_match.group(1).strip() if notes_match else ""
            title, blocks = parse_blocks(NOTES_RE.sub("", raw))
            if not title and not blocks:
                continue
            try:
                deck.content(title, blocks, notes, lesson["title"])
            except UnknownMacro as exc:
                # Locate the offending span for the author rather than dumping
                # a traceback from inside the transliterator.
                raise SystemExit(
                    f"{lesson.get('file', lesson['title'])}: slide {n} "
                    f"({title or 'untitled'}): {exc}"
                ) from None

    out_path = os.path.join(out_dir, f"{module.get('slug') or _slug(module['title'])}.pptx")
    deck.save(out_path)

    if deck.missing_assets:
        msg = f"{out_path}: {len(deck.missing_assets)} unresolved image(s): " + ", ".join(
            sorted(set(deck.missing_assets))[:6]
        )
        if strict_assets:
            raise SystemExit(msg)
        print(f"  warning: {msg}", file=sys.stderr)

    return out_path, len(deck.prs.slides._sldIdLst)


def read_lesson_body(base: str, lesson: dict) -> str:
    if lesson.get("file"):
        path = os.path.join(base, lesson["file"])
        if not os.path.isfile(path):
            raise SystemExit(f"lesson file not found: {path}")
        with open(path, encoding="utf-8") as fh:
            return fh.read()
    if "body_md" in lesson:
        return lesson["body_md"]
    raise SystemExit(f"lesson {lesson.get('title')!r} has neither 'file' nor 'body_md'")


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("course_dir", help="directory containing course.yaml")
    ap.add_argument("--out", default="build/slides", help="output directory")
    ap.add_argument("--module", action="append", help="build only this module slug (repeatable)")
    ap.add_argument("--strict-assets", action="store_true",
                    help="fail instead of warning when an image cannot be resolved")
    args = ap.parse_args()

    base = os.path.abspath(args.course_dir)
    manifest = load_manifest(base)
    course = manifest.get("course") or {}
    modules = manifest.get("modules") or []
    if args.module:
        wanted = set(args.module)
        modules = [m for m in modules if (m.get("slug") or _slug(m["title"])) in wanted]
        if not modules:
            raise SystemExit(f"no module matched {sorted(wanted)}")

    os.makedirs(args.out, exist_ok=True)
    math = MathRenderer(os.path.join(args.out, ".mathcache"))
    for module in modules:
        path, count = build_module(base, course, module, args.out,
                                   args.strict_assets, math)
        print(f"{path}  ({count} slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
